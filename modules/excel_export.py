import os
import pandas as pd
from pathlib import Path

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
_OUTPUT_DIR   = _PROJECT_ROOT / "static" / "outputs"


try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side)
    from openpyxl.utils import get_column_letter
    OPENPYXL = True
except ImportError:
    OPENPYXL = False


def _fill(color):
    return PatternFill("solid", fgColor=color)

def _font(color="FFFFFF", bold=True, size=11):
    return Font(color=color, bold=bold, size=size)

def _center():
    return Alignment(horizontal="center",
                     vertical="center", wrap_text=True)

def _border():
    s = Side(style="thin", color="DDDDDD")
    return Border(left=s, right=s, top=s, bottom=s)

def _col_w(ws, col, width):
    ws.column_dimensions[
        get_column_letter(col)].width = width


def generate_excel(filename, df_path,
                   eda_stats, sentiment_stats, insights):
    if not OPENPYXL:
        return None

    _OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = str(_OUTPUT_DIR / "lumiq_report.xlsx")


    try:
        df = pd.read_csv(df_path)
    except Exception:
        df = pd.DataFrame()

    wb  = Workbook()
    PUR = "534AB7"
    GRN = "1D9E75"
    RED = "E24B4A"
    LGT = "EEEDFE"
    WHT = "FFFFFF"
    GRY = "F5F5F7"
    DRK = "1A1A2E"

    # ── SHEET 1: Overview ──────────────────────────────────
    ws1 = wb.active
    ws1.title = "Overview"
    ws1.sheet_view.showGridLines = False

    ws1.merge_cells("A1:F1")
    ws1["A1"] = "LUMIQ - AI Data Analysis Report"
    ws1["A1"].font      = _font(size=16)
    ws1["A1"].fill      = _fill(PUR)
    ws1["A1"].alignment = _center()
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:F2")
    ws1["A2"] = f"File: {filename}"
    ws1["A2"].font      = Font(color=PUR, size=11)
    ws1["A2"].alignment = _center()
    ws1.row_dimensions[2].height = 22

    ws1.row_dimensions[3].height = 10

    kpis = [
        ("Total Rows",
         f'{eda_stats.get("total_rows") or 0:,}', PUR),
        ("Columns",
         str(eda_stats.get("total_cols") or 0), PUR),
        ("Missing Data",
         f'{eda_stats.get("missing_pct") or 0}%', RED),
        ("Dupes Removed",
         str(eda_stats.get("dupes_removed") or 0), RED),
        ("Positive %",
         f'{sentiment_stats.get("positive_pct") or 0}%', GRN),
        ("Best Accuracy",
         f'{(sentiment_stats.get("ml") or {}).get("best_accuracy", 0) or 0}%',
         GRN),
    ]

    for i, (label, value, color) in enumerate(kpis):
        col = i + 1
        lc = ws1.cell(row=4, column=col, value=label)
        lc.font      = Font(bold=True, size=9,
                            color="666666")
        lc.alignment = _center()
        ws1.row_dimensions[4].height = 18

        vc = ws1.cell(row=5, column=col, value=value)
        vc.font      = Font(bold=True, size=18,
                            color=color)
        vc.fill      = _fill(GRY)
        vc.alignment = _center()
        ws1.row_dimensions[5].height = 44
        _col_w(ws1, col, 22)

    ws1.row_dimensions[6].height = 14

    ws1.cell(row=7, column=1,
             value="AI-Generated Insights").font = \
        Font(bold=True, size=12, color=PUR)
    ws1.row_dimensions[7].height = 22

    row = 8
    for ins in insights:
        ws1.merge_cells(f"A{row}:F{row}")
        c = ws1.cell(
            row=row, column=1,
            value=f"* {ins.get('title','')}: "
                  f"{ins.get('text','')}")
        c.font      = Font(size=10, color=DRK)
        c.alignment = Alignment(wrap_text=True,
                                vertical="top")
        ws1.row_dimensions[row].height = 30
        row += 1

    # ── SHEET 2: Cleaned Data ──────────────────────────────
    if not df.empty:
        ws2 = wb.create_sheet("Cleaned Data")
        ws2.sheet_view.showGridLines = False

        for ci, cn in enumerate(df.columns, 1):
            c = ws2.cell(row=1, column=ci,
                         value=str(cn))
            c.font      = _font()
            c.fill      = _fill(PUR)
            c.alignment = _center()
            c.border    = _border()
            _col_w(ws2, ci, 18)
        ws2.row_dimensions[1].height = 24

        max_rows = min(len(df), 1000)
        for ri in range(max_rows):
            for ci, cn in enumerate(df.columns, 1):
                val = df.iloc[ri, ci-1]
                if pd.isna(val):
                    val = ""
                elif isinstance(val, float):
                    val = round(val, 3)
                c = ws2.cell(row=ri+2, column=ci,
                             value=val)
                c.border    = _border()
                c.alignment = Alignment(
                    vertical="center")
                if ri % 2 == 0:
                    c.fill = _fill(GRY)

    # ── SHEET 3: EDA Summary ───────────────────────────────
    ws3 = wb.create_sheet("EDA Summary")
    ws3.sheet_view.showGridLines = False

    ws3.merge_cells("A1:G1")
    ws3["A1"] = "Exploratory Data Analysis Summary"
    ws3["A1"].font      = _font(size=13)
    ws3["A1"].fill      = _fill(PUR)
    ws3["A1"].alignment = _center()
    ws3.row_dimensions[1].height = 30

    num_summary = eda_stats.get("numeric_summary") or {}
    if num_summary:
        headers = ["Column","Mean","Median",
                   "Std Dev","Min","Max","Outliers"]
        for ci, h in enumerate(headers, 1):
            c = ws3.cell(row=3, column=ci, value=h)
            c.font      = _font(size=10)
            c.fill      = _fill(PUR)
            c.alignment = _center()
            c.border    = _border()
            _col_w(ws3, ci, 14)

        for i, (cn, s) in enumerate(
                num_summary.items()):
            row = i + 4
            vals = [
                cn,
                s.get("mean",   0),
                s.get("median", 0),
                s.get("std",    0),
                s.get("min",    0),
                s.get("max",    0),
                s.get("outliers",0),
            ]
            for ci, val in enumerate(vals, 1):
                c = ws3.cell(row=row, column=ci,
                             value=val)
                c.border    = _border()
                c.alignment = _center()
                if i % 2 == 0:
                    c.fill = _fill(GRY)
                if ci == 7 and val > 0:
                    c.font = Font(bold=True, color=RED)

    # ── SHEET 4: Sentiment ─────────────────────────────────
    ws4 = wb.create_sheet("Sentiment Analysis")
    ws4.sheet_view.showGridLines = False

    ws4.merge_cells("A1:D1")
    ws4["A1"] = "Sentiment Analysis Results"
    ws4["A1"].font      = _font(size=13)
    ws4["A1"].fill      = _fill(PUR)
    ws4["A1"].alignment = _center()
    ws4.row_dimensions[1].height = 30

    if sentiment_stats.get("available"):
        pairs = [
            ("Text Column",
             sentiment_stats.get("text_column","N/A")),
            ("Total Samples",
             sentiment_stats.get("total_samples",0)),
            ("Positive Count",
             sentiment_stats.get("positive",0)),
            ("Positive %",
             f'{sentiment_stats.get("positive_pct",0)}%'),
            ("Negative Count",
             sentiment_stats.get("negative",0)),
            ("Negative %",
             f'{sentiment_stats.get("negative_pct",0)}%'),
            ("Neutral Count",
             sentiment_stats.get("neutral",0)),
            ("Neutral %",
             f'{sentiment_stats.get("neutral_pct",0)}%'),
        ]
        for i, (lbl, val) in enumerate(pairs):
            row = i + 3
            lc = ws4.cell(row=row, column=1, value=lbl)
            lc.font   = Font(bold=True, size=10)
            lc.fill   = _fill(LGT)
            lc.border = _border()
            _col_w(ws4, 1, 22)

            vc = ws4.cell(row=row, column=2, value=val)
            vc.border    = _border()
            vc.alignment = _center()
            _col_w(ws4, 2, 18)

        ml = sentiment_stats.get("ml") or {}
        if ml.get("available"):
            sr = len(pairs) + 5
            ws4.cell(row=sr-1, column=1,
                     value="ML Model Results").font = \
                Font(bold=True, size=12, color=PUR)

            mh = ["Model","Accuracy","Precision",
                  "Recall","F1","CV Mean","Time(s)"]
            for ci, h in enumerate(mh, 1):
                c = ws4.cell(row=sr, column=ci, value=h)
                c.font      = _font(size=10)
                c.fill      = _fill(PUR)
                c.alignment = _center()
                c.border    = _border()
                _col_w(ws4, ci, 15)

            for i, (mn, mr) in enumerate(
                    ml.get("model_results",{}).items()):
                row  = sr + i + 1
                best = (mn == ml.get("best_model"))
                vals = [
                    mn,
                    f'{mr.get("accuracy","N/A")}%',
                    f'{mr.get("precision","N/A")}%',
                    f'{mr.get("recall","N/A")}%',
                    f'{mr.get("f1","N/A")}%',
                    f'{mr.get("cv_mean","N/A")}%',
                    mr.get("train_time_sec","N/A"),
                ]
                for ci, val in enumerate(vals, 1):
                    c = ws4.cell(row=row, column=ci,
                                 value=val)
                    c.border    = _border()
                    c.alignment = _center()
                    if best:
                        c.fill = _fill(LGT)
                        c.font = Font(bold=True)

    # ── SHEET 5: Keywords ──────────────────────────────────
    ws5 = wb.create_sheet("Keywords")
    ws5.sheet_view.showGridLines = False

    ws5.merge_cells("A1:F1")
    ws5["A1"] = "Top Keywords and Sentiment Words"
    ws5["A1"].font      = _font(size=13)
    ws5["A1"].fill      = _fill(PUR)
    ws5["A1"].alignment = _center()
    ws5.row_dimensions[1].height = 30

    sections = [
        ("All Keywords",
         sentiment_stats.get("keywords") or [], 1),
        ("Positive Keywords",
         sentiment_stats.get("pos_keywords") or [], 3),
        ("Negative Keywords",
         sentiment_stats.get("neg_keywords") or [], 5),
    ]

    for title, kws, sc in sections:
        ws5.cell(row=3, column=sc,
                 value=title).font = \
            Font(bold=True, size=11, color=PUR)
        for i, kw in enumerate(kws[:15]):
            row = i + 4
            wc  = ws5.cell(row=row, column=sc,
                           value=kw.get("word",""))
            wc.border    = _border()
            wc.alignment = _center()
            if i % 2 == 0:
                wc.fill = _fill(GRY)

            sc2 = ws5.cell(
                row=row, column=sc+1,
                value=round(kw.get("score",0)*100,1))
            sc2.border    = _border()
            sc2.alignment = _center()
            if i % 2 == 0:
                sc2.fill = _fill(GRY)

        _col_w(ws5, sc,   18)
        _col_w(ws5, sc+1, 12)

    wb.save(out_path)
    return out_path